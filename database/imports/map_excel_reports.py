"""Extrai resultados estruturados de planilhas laboratoriais sem gravar no banco.

Gera um JSON de prévia para revisão. A carga PostgreSQL deve consumir somente
registros com ``review_required == false`` após a conferência clínica.
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


COPY_SUFFIX = re.compile(r"(?:\s*\(\d+\))+$")
RECORD = re.compile(r"^(?P<prefix>[A-Za-z]?)(?P<record>\d+)_26\s*(?P<kind>.*)$", re.I)


def canonical_kind(file_name: str) -> str:
    base = COPY_SUFFIX.sub("", Path(file_name).stem).upper().strip()
    match = RECORD.match(base)
    kind = (match.group("kind") if match else base).strip()
    compact = re.sub(r"[^A-Z0-9]", "", unicodedata.normalize("NFD", kind).encode("ascii", "ignore").decode().upper())
    if compact in {"B", "BCOLTRIG"}:
        return "BIOQUIMICA"
    if compact == "N":
        return "BIOQUIMICA"
    if compact.startswith("EAS"):
        return "URINALISE"
    if compact in {"4DX", "IMUNO4DX"}:
        return "IMUNO_4DX"
    aliases = {
        "KNOTT": "KNOTT", "KNOOT": "KNOTT", "LEPTO": "LEPTOSPIROSE",
        "COMPAT": "COMPATIBILIDADE", "COMPATIBILIDADE": "COMPATIBILIDADE",
        "HEMOAGLUTINACAO": "HEMOAGLUTINACAO", "CORTISOL": "CORTISOL",
        "T4": "T4", "LAVADOBRONCOALVEOLAR": "LAVADO_BRONCOALVEOLAR",
        "LIQCAV": "LIQUIDO_CAVITARIO", "LC": "LIQUIDO_CAVITARIO",
        "PESQUISADEHEMOPARASITA": "PESQUISA_HEMOPARASITA",
        "PONTAORELHA": "PONTA_ORELHA", "SALINA": "SALINA",
    }
    if compact in aliases:
        return aliases[compact]
    if compact.startswith("H"):
        return "HEMOGRAMA"
    return "NAO_CLASSIFICADO"


def source_record(file_name: str) -> str | None:
    base = COPY_SUFFIX.sub("", Path(file_name).stem)
    match = RECORD.match(base)
    return match.group("record") if match else None


def source_origin(file_name: str) -> str | None:
    """Arquivos iniciados por A pertencem ao Projeto AMA; os demais, ao HV."""
    base = COPY_SUFFIX.sub("", Path(file_name).stem)
    match = RECORD.match(base)
    if not match:
        return None
    return "AMA" if match.group("prefix").upper() == "A" else "HV"


def is_copy_file(path: Path) -> bool:
    return bool(COPY_SUFFIX.search(path.stem))


def label_and_unit(value: object) -> tuple[str, str | None]:
    text = str(value or "").strip()
    match = re.match(r"^(.*?)\s*\(([^)]+)\)$", text)
    return (match.group(1).strip(), match.group(2).strip()) if match else (text, None)


def numeric(value: object) -> float | None:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def analyte_code(name: str, unit: str | None) -> str:
    base = unicodedata.normalize("NFD", name).encode("ascii", "ignore").decode().upper()
    base = re.sub(r"[^A-Z0-9]+", "_", base).strip("_")
    normalized_unit = unicodedata.normalize("NFD", (unit or "").replace("µ", "U")).encode("ascii", "ignore").decode().upper()
    if normalized_unit == "%": return f"{base}_PERCENTUAL"
    if "UL" in normalized_unit: return f"{base}_ABSOLUTO_UL"
    return base


def result(section: str, label: object, value: object, low: object = None, high: object = None, reference: object = None, unit: str | None = None) -> dict | None:
    name, label_unit = label_and_unit(label)
    result_number = numeric(value)
    result_text = None if result_number is not None else (str(value).strip() if value not in (None, "") else None)
    if not name or (result_number is None and not result_text):
        return None
    effective_unit = unit or label_unit
    return {"section": section, "analyte_code": analyte_code(name, effective_unit), "analyte_name": name, "unit": effective_unit,
            "result_numeric": result_number, "result_text": result_text,
            "reference_low": numeric(low), "reference_high": numeric(high),
            "reference_text": str(reference).strip() if reference not in (None, "") else None}


def hemogram(sheet) -> list[dict]:
    rows: list[dict] = []
    for row in range(15, 21):
        item = result("Eritrograma", sheet.cell(row, 1).value, sheet.cell(row, 5).value, sheet.cell(row, 10).value, sheet.cell(row, 12).value)
        if item: rows.append(item)
    rows.append(result("Leucograma", "Leucócitos (/µL)", sheet.cell(23, 5).value, sheet.cell(23, 10).value, sheet.cell(23, 12).value, unit="/µL"))
    for row in range(24, 32):
        label = sheet.cell(row, 1).value
        percent = result("Leucograma", f"{label} (%)", sheet.cell(row, 4).value, sheet.cell(row, 8).value, sheet.cell(row, 9).value, unit="%")
        absolute = result("Leucograma", f"{label} (/µL)", sheet.cell(row, 5).value, sheet.cell(row, 10).value, sheet.cell(row, 12).value, unit="/µL")
        if percent: rows.append(percent)
        if absolute: rows.append(absolute)
    for row in (33, 34):
        item = result("Hemograma", sheet.cell(row, 1).value, sheet.cell(row, 5).value, sheet.cell(row, 10).value, sheet.cell(row, 12).value)
        if item: rows.append(item)
    item = result("Contagem de reticulócitos", sheet.cell(37, 1).value, sheet.cell(37, 4).value, reference=sheet.cell(37, 9).value)
    return rows + ([item] if item else [])


def biochemistry(sheet) -> list[dict]:
    rows = []
    for row in range(14, 80):
        label = sheet.cell(row, 1).value
        if label and "automatizado" not in str(label).lower():
            item = result("Bioquímica", label, sheet.cell(row, 6).value, sheet.cell(row, 10).value, sheet.cell(row, 12).value)
            if item: rows.append(item)
    return rows


def urinalysis(sheet) -> list[dict]:
    rows, section = [], "Urinálise"
    for row in range(12, 80):
        heading = str(sheet.cell(row, 1).value or "").strip().lower()
        if heading in {"exame físico", "exame quimico", "exame químico", "sedimentoscopia", "bioquímica urinária"}:
            section = heading.title()
            continue
        label = sheet.cell(row, 1).value
        if not label or str(label).lower().startswith(("observa", "médicos")):
            continue
        value = next((sheet.cell(row, col).value for col in (5, 4, 3)
                      if sheet.cell(row, col).value not in (None, "")), None)
        item = result(section, label, value, sheet.cell(row, 10).value, sheet.cell(row, 12).value, sheet.cell(row, 10).value)
        if item: rows.append(item)
    return rows


def qualitative(sheet, kind: str) -> list[dict]:
    """Mapeia painéis qualitativos cujo resultado ocupa as colunas B/E."""
    rows, section = [], kind.replace("_", " ").title()
    for row in range(10, min(sheet.max_row, 80) + 1):
        left = str(sheet.cell(row, 1).value or "").strip()
        label = str(sheet.cell(row, 2).value or "").strip()
        value = sheet.cell(row, 5).value
        if any(token in left.casefold() for token in ("observa", "médicos veterinários")):
            break
        if left.casefold().startswith("resultado") and label and value not in (None, ""):
            item = result(section, label, value)
        elif label and value not in (None, "") and row <= 25:
            item = result(section, label, value)
        elif left.casefold().startswith("resultado") and value not in (None, ""):
            item = result(section, f"{section} — Resultado", value)
        else:
            item = None
        if item:
            rows.append(item)
    return rows


def cortisol(sheet) -> list[dict]:
    rows = []
    for row in range(10, min(sheet.max_row, 30) + 1):
        label, value, unit = sheet.cell(row, 1).value, sheet.cell(row, 3).value, sheet.cell(row, 4).value
        if label and value not in (None, "") and str(label).strip().casefold() not in {"resultado:", "interpretação do teste:"}:
            item = result("Cortisol", label, value, unit=str(unit or "").strip() or None)
            if item:
                rows.append(item)
    return rows


def t4(sheet) -> list[dict]:
    for row in range(10, min(sheet.max_row, 25) + 1):
        value, unit = sheet.cell(row, 3).value, sheet.cell(row, 4).value
        if value not in (None, ""):
            item = result("T4", "T4", value, unit=str(unit or "").strip() or None)
            return [item] if item else []
    return []


def liquid_analysis(sheet, kind: str) -> list[dict]:
    rows, section = [], "Exame físico"
    for row in range(15, min(sheet.max_row, 80) + 1):
        label = sheet.cell(row, 1).value
        normalized = str(label or "").strip().casefold()
        if "citologia" in normalized or normalized.startswith(("descrição", "diagnóstico", "comentário", "médicos")):
            break
        if normalized in {"exame físico", "exame quimico", "exame químico", "exame bioquimico", "exame bioquímico"}:
            section = str(label).strip().title()
            continue
        if not label or any(term in normalized for term in ("colorimétrico", "diferença dos valores", "resultados")):
            continue
        if section.startswith("Exame Bioqu"):
            for column, suffix in ((6, "Efusão"), (10, "Soro")):
                name, unit = label_and_unit(label)
                item = result(section, f"{name} — {suffix}", sheet.cell(row, column).value, unit=unit)
                if item:
                    rows.append(item)
        else:
            unit = sheet.cell(row, 8).value if row == 19 else None
            item = result(section, label, sheet.cell(row, 5).value, unit=str(unit or "").strip() or None)
            if item:
                rows.append(item)
    return rows


def compatibility(sheet) -> list[dict]:
    rows = []
    for row in (13, 15, 17):
        subject = str(sheet.cell(row, 1).value or "").strip()
        if not subject:
            continue
        for label, column in (("Hematócrito (%)", 3), ("PPT (g/dL)", 8)):
            item = result("Compatibilidade sanguínea", f"{subject} — {label}", sheet.cell(row, column).value)
            if item:
                rows.append(item)
    for row in (24, 26, 28, 30):
        label = str(sheet.cell(row, 3).value or "").strip()
        if not label:
            continue
        for subject, column in (("Doador 1", 6), ("Doador 2", 9)):
            item = result("Compatibilidade sanguínea", f"{label} — {subject}", sheet.cell(row, column).value)
            if item:
                rows.append(item)
    return rows


def infer_kind(sheet, initial_kind: str) -> str:
    if initial_kind != "NAO_CLASSIFICADO":
        return initial_kind
    text = "\n".join(" ".join(str(value or "") for value in row[:10]) for row in sheet.iter_rows(min_row=1, max_row=30, values_only=True)).casefold()
    if "hemograma" in text:
        return "HEMOGRAMA"
    if "bioquímica" in text or "bioquimica" in text:
        return "BIOQUIMICA"
    return initial_kind


def generic(sheet, kind: str) -> list[dict]:
    """Extrai tabelas simples e sempre exige revisão clínica antes da carga."""
    rows, section = [], kind.replace("_", " ").title()
    for row in range(1, min(sheet.max_row, 120) + 1):
        label = sheet.cell(row, 1).value
        if not label or len(str(label)) > 80:
            continue
        value = next((sheet.cell(row, col).value for col in range(2, min(sheet.max_column, 12) + 1) if sheet.cell(row, col).value not in (None, "")), None)
        item = result(section, label, value)
        if item: rows.append(item)
    return rows


def worksheet_text(sheet) -> str:
    """Preserva o texto integral disponível na planilha para auditoria e pesquisa."""
    lines = []
    for row in sheet.iter_rows(values_only=True):
        values = [str(value).strip() for value in row if value not in (None, "")]
        if values:
            lines.append(" | ".join(values))
    return "\n".join(lines)


def comments_text(sheet) -> str:
    """Extrai, sem resumir, o bloco de comentários/observações até a assinatura."""
    comments, collecting = [], False
    for row in sheet.iter_rows(values_only=True):
        values = [str(value).strip() for value in row if value not in (None, "")]
        if not values:
            continue
        line = " | ".join(values)
        normalized = unicodedata.normalize("NFD", line).encode("ascii", "ignore").decode().casefold()
        if "medicos veterinarios responsaveis" in normalized:
            break
        if not collecting and re.search(r"\b(comentarios?|observacoes?)\s*:", normalized):
            collecting = True
        if collecting:
            comments.append(line)
    return "\n".join(comments)


def map_file(path: Path) -> dict:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    kind = infer_kind(sheet, canonical_kind(path.name))
    mapper = {
        "HEMOGRAMA": hemogram, "BIOQUIMICA": biochemistry, "URINALISE": urinalysis,
        "IMUNO_4DX": qualitative, "LEPTOSPIROSE": qualitative, "KNOTT": qualitative,
        "HEMOAGLUTINACAO": qualitative, "PESQUISA_HEMOPARASITA": qualitative,
        "PONTA_ORELHA": qualitative, "SALINA": qualitative,
        "CORTISOL": cortisol, "T4": t4, "COMPATIBILIDADE": compatibility,
        "LIQUIDO_CAVITARIO": liquid_analysis, "LAVADO_BRONCOALVEOLAR": liquid_analysis,
    }.get(kind)
    if mapper in {qualitative, liquid_analysis}:
        values = mapper(sheet, kind)
    else:
        values = mapper(sheet) if mapper else generic(sheet, kind)
    return {"source_file": path.name, "source_record": source_record(path.name), "exam_type": kind,
            "source_origin": source_origin(path.name), "worksheet": sheet.title,
            "results": [item for item in values if item], "raw_text": worksheet_text(sheet),
            "comments": comments_text(sheet), "review_required": mapper is None}


def main(directory: Path, output: Path) -> None:
    files = sorted(path for path in directory.glob("*.xlsx") if not is_copy_file(path))
    mapped = [map_file(path) for path in files]
    output.write_text(json.dumps({"records": mapped}, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Tipos encontrados:")
    for kind, count in sorted(Counter(item["exam_type"] for item in mapped).items()): print(f"- {kind}: {count}")
    print(f"Prévia gerada: {output}")


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
