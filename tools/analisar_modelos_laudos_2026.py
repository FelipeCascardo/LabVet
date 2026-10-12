from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Projetos\LabVet\Importação\Laudos 2026")
OUTPUT = Path(r"C:\Projetos\LabVet\tmp\analise_modelos_laudos_2026.json")


def normalizar(valor: object) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(caractere for caractere in texto if not unicodedata.combining(caractere))
    return re.sub(r"\s+", " ", texto).strip().casefold()


def texto(valor: object) -> str:
    return "" if valor is None else re.sub(r"\s+", " ", str(valor)).strip()


def grupo_animal(caminho: Path) -> str:
    partes = caminho.relative_to(ROOT).parts[:-1]
    return " > ".join(re.sub(r"\s*2026\s*", "", parte, flags=re.I).strip() for parte in partes)


def extrair_idade(nome: str, linhas: list[list[str]]) -> list[str]:
    linhas_idade = [
        " ".join(linha)
        for linha in linhas[:100]
        if any(marcador in normalizar(" ".join(linha)) for marcador in ("idade", "faixa etaria"))
    ]
    corpus = " ".join([nome] + linhas_idade)
    padrao = re.compile(
        r"(?:filhote|adulto|adultos|jovem|idoso|neonato|menos\s+de\s+\d+\s*m[eê]s|"
        r"at[eé]\s+\d+\s*(?:m[eê]s|meses|ano|anos)|\d+\s*(?:m[eê]s|meses|ano|anos))",
        re.I,
    )
    return list(dict.fromkeys(re.sub(r"\s+", " ", trecho).strip() for trecho in padrao.findall(corpus)))


def analisar_arquivo(caminho: Path) -> dict:
    livro = load_workbook(caminho, read_only=True, data_only=False)
    planilhas = []
    for planilha in livro.worksheets:
        linhas = []
        for linha in planilha.iter_rows(values_only=True):
            valores = [texto(valor) for valor in linha]
            if any(valores):
                linhas.append(valores)
        planilhas.append(
            {
                "nome": planilha.title,
                "linhas": linhas[:250],
                "linhasTotal": len(linhas),
                "colunasTotal": planilha.max_column,
            }
        )

    todas_linhas = [linha for planilha in planilhas for linha in planilha["linhas"]]
    cabecalhos = []
    for linha in todas_linhas:
        combinada = " | ".join(linha)
        normalizada = normalizar(combinada)
        if any(palavra in normalizada for palavra in ("resultado", "referencia", "unidade", "analito", "exame")):
            cabecalhos.append(combinada)

    return {
        "arquivo": str(caminho.relative_to(ROOT)),
        "grupoAnimal": grupo_animal(caminho),
        "idadeMencionada": extrair_idade(caminho.stem, todas_linhas),
        "planilhas": planilhas,
        "linhasCabecalho": list(dict.fromkeys(cabecalhos))[:30],
    }


def main() -> None:
    resultados = []
    erros = []
    for caminho in sorted(ROOT.rglob("*.xlsx")):
        try:
            resultados.append(analisar_arquivo(caminho))
        except Exception as erro:  # relatório de arquivos que precisam de revisão manual
            erros.append({"arquivo": str(caminho.relative_to(ROOT)), "erro": str(erro)})

    resumo = Counter(item["grupoAnimal"] for item in resultados)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(
        json.dumps(
            {
                "quantidadeXlsx": len(resultados),
                "quantidadeXlsNaoLidos": len(list(ROOT.rglob("*.xls"))),
                "gruposAnimal": dict(sorted(resumo.items())),
                "modelos": resultados,
                "erros": erros,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(OUTPUT)
    print(f"xlsx={len(resultados)} xls_nao_lidos={len(list(ROOT.rglob('*.xls')))} erros={len(erros)}")


if __name__ == "__main__":
    main()
